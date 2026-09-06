from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "artifacts"
SCREENSHOT_DIR = OUT_DIR / "screenshots"
OUT_FILE = OUT_DIR / "shenzhen-sci-tech-platform-client-intro.pptx"


SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Noto Sans CJK SC"
FONT_LATIN = "Aptos"

NAVY = RGBColor(20, 61, 89)
NAVY_DARK = RGBColor(13, 44, 65)
INK = RGBColor(24, 38, 50)
MUTED = RGBColor(101, 119, 130)
TEAL = RGBColor(22, 139, 130)
TEAL_SOFT = RGBColor(231, 245, 242)
BLUE_SOFT = RGBColor(237, 244, 248)
AMBER = RGBColor(190, 123, 29)
AMBER_SOFT = RGBColor(255, 244, 223)
CANVAS = RGBColor(247, 249, 252)
SURFACE = RGBColor(255, 255, 255)
LINE = RGBColor(218, 227, 232)
PALE = RGBColor(242, 246, 248)


def inch(value: float) -> Inches:
    return Inches(value)


def add_shape(slide, shape_type, x, y, w, h, *, fill=None, line=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, inch(x), inch(y), inch(w), inch(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.8)
    return shape


def rounded(slide, x, y, w, h, *, fill=SURFACE, line=LINE):
    return add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=fill,
        line=line,
    )


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=12,
    color=INK,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    fit=False,
):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = inch(margin)
    box.text_frame.margin_right = inch(margin)
    box.text_frame.margin_top = inch(margin)
    box.text_frame.margin_bottom = inch(margin)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if fit:
        box.text_frame.fit_text(font_family=font, max_size=Pt(size))
    return box


def add_rich_text(slide, runs, x, y, w, h, *, size=12, color=INK, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = inch(0.04)
    box.text_frame.margin_right = inch(0.04)
    box.text_frame.margin_top = inch(0.04)
    box.text_frame.margin_bottom = inch(0.04)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    for text, kwargs in runs:
        run = p.add_run()
        run.text = text
        run.font.name = kwargs.get("font", FONT)
        run.font.size = Pt(kwargs.get("size", size))
        run.font.bold = kwargs.get("bold", False)
        run.font.color.rgb = kwargs.get("color", color)
    return box


def add_rule(slide, x, y, w, *, color=LINE, width=1.0):
    line = slide.shapes.add_connector(
        1,
        inch(x),
        inch(y),
        inch(x + w),
        inch(y),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_arrow(slide, x1, y1, x2, y2, *, color=TEAL, width=1.7):
    line = slide.shapes.add_connector(1, inch(x1), inch(y1), inch(x2), inch(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def page_header(slide, kicker, title, subtitle, page):
    add_text(slide, kicker.upper(), 0.58, 0.28, 7.8, 0.22, size=8.5, color=TEAL, bold=True)
    add_text(slide, title, 0.58, 0.57, 11.45, 0.55, size=24, color=NAVY_DARK, bold=True)
    add_text(slide, subtitle, 0.60, 1.18, 10.9, 0.3, size=11.2, color=MUTED)
    add_text(slide, f"0{page} / 02", 12.05, 0.35, 0.72, 0.24, size=8.5, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    add_rule(slide, 0.58, 1.58, 12.16, color=LINE, width=0.9)


def screenshot_card(slide, image_path, x, y, w, h, caption, *, accent=TEAL):
    rounded(slide, x, y, w, h, fill=SURFACE, line=LINE)
    slide.shapes.add_picture(str(image_path), inch(x + 0.09), inch(y + 0.09), width=inch(w - 0.18), height=inch(h - 0.38))
    add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.09, y + 0.09, 0.06, h - 0.38, fill=accent, line=accent)
    add_text(slide, caption, x + 0.12, y + h - 0.25, w - 0.24, 0.16, size=8.4, color=MUTED, bold=True)


def chip(slide, text, x, y, w, *, fill=TEAL_SOFT, color=TEAL, line=None, size=10.2, bold=True):
    rounded(slide, x, y, w, 0.36, fill=fill, line=line or fill)
    add_text(slide, text, x, y + 0.01, w, 0.31, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def stage_box(slide, number, title, note, x, y, w, *, fill=SURFACE, accent=TEAL):
    rounded(slide, x, y, w, 1.08, fill=fill, line=LINE)
    add_shape(slide, MSO_SHAPE.OVAL, x + 0.12, y + 0.14, 0.28, 0.28, fill=accent, line=accent)
    add_text(slide, str(number), x + 0.12, y + 0.145, 0.28, 0.23, size=8.2, color=SURFACE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, title, x + 0.48, y + 0.14, w - 0.6, 0.23, size=10.3, color=NAVY_DARK, bold=True)
    add_text(slide, note, x + 0.12, y + 0.54, w - 0.24, 0.34, size=8.7, color=MUTED)


def layer_band(slide, y, h, label, label_color, cards):
    rounded(slide, 0.58, y, 8.04, h, fill=SURFACE, line=LINE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.58, y, 0.11, h, fill=label_color, line=label_color)
    add_text(slide, label, 0.83, y + 0.16, 1.44, 0.28, size=10.7, color=label_color, bold=True)
    card_x = 2.25
    gap = 0.1
    total = 8.04 - (card_x - 0.58) - 0.16
    card_w = (total - gap * (len(cards) - 1)) / len(cards)
    for idx, (title, note) in enumerate(cards):
        x = card_x + idx * (card_w + gap)
        rounded(slide, x, y + 0.16, card_w, h - 0.32, fill=PALE, line=PALE)
        add_text(slide, title, x + 0.08, y + 0.26, card_w - 0.16, 0.22, size=9.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, note, x + 0.08, y + 0.53, card_w - 0.16, h - 0.62, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)


def evidence_item(slide, n, title, note, y, *, color=TEAL):
    add_shape(slide, MSO_SHAPE.OVAL, 9.12, y + 0.02, 0.27, 0.27, fill=color, line=color)
    add_text(slide, str(n), 9.12, y + 0.025, 0.27, 0.22, size=8, color=SURFACE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, title, 9.52, y, 2.95, 0.22, size=9.7, color=SURFACE, bold=True)
    add_text(slide, note, 9.52, y + 0.24, 2.95, 0.34, size=8.1, color=RGBColor(190, 211, 219))


def build_deck():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    home = SCREENSHOT_DIR / "baseline-home.png"
    professors = SCREENSHOT_DIR / "baseline-professors-list.png"
    if not home.exists() or not professors.exists():
        raise FileNotFoundError("Expected real system screenshots under artifacts/screenshots/")

    prs = Presentation()
    prs.slide_width = inch(SLIDE_W)
    prs.slide_height = inch(SLIDE_H)
    blank = prs.slide_layouts[6]

    # Slide 1: user-facing value and retrieval path.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CANVAS
    page_header(
        slide,
        "Shenzhen Sci-Tech Data Platform",
        "从一句话问题，到可追溯的科创答案",
        "面向科创生态的对话式检索入口：用户只需提问，系统负责理解、路由、检索与组织答案。",
        1,
    )

    add_text(slide, "客户看到的是答案，系统完成的是一条可解释的检索链路", 0.6, 1.82, 7.55, 0.28, size=12.2, color=NAVY_DARK, bold=True)
    add_text(slide, "单域查询、跨域聚合、多轮追问，共享同一套证据与身份约束。", 0.6, 2.12, 7.4, 0.24, size=9.6, color=MUTED)

    stages = [
        ("问题理解", "识别对象、条件与时间"),
        ("意图路由", "决定查哪个域 / 是否跨域"),
        ("检索执行", "语义召回 + 结构化筛选"),
        ("融合排序", "多源召回、融合、rerank"),
        ("答案组织", "结构化回答 + evidence"),
    ]
    xs = [0.6, 2.12, 3.64, 5.16, 6.68]
    for i, ((title, note), x) in enumerate(zip(stages, xs, strict=True), start=1):
        fill = TEAL_SOFT if i in {1, 5} else SURFACE
        stage_box(slide, i, title, note, x, 2.53, 1.28, fill=fill, accent=TEAL if i != 3 else AMBER)
        if i < 5:
            add_arrow(slide, x + 1.31, 3.07, x + 1.48, 3.07, color=TEAL, width=1.2)

    add_text(slide, "四类科创数据域", 0.6, 3.97, 1.45, 0.23, size=9.3, color=MUTED, bold=True)
    chip(slide, "教授", 2.05, 3.9, 1.15)
    chip(slide, "企业", 3.34, 3.9, 1.15, fill=BLUE_SOFT, color=NAVY)
    chip(slide, "论文", 4.63, 3.9, 1.15, fill=AMBER_SOFT, color=AMBER)
    chip(slide, "专利", 5.92, 3.9, 1.15, fill=PALE, color=NAVY)

    rounded(slide, 0.6, 4.55, 7.36, 1.45, fill=NAVY_DARK, line=NAVY_DARK)
    add_text(slide, "三个让客户放心的体验", 0.86, 4.79, 2.5, 0.24, size=10.5, color=RGBColor(182, 220, 214), bold=True)
    bullets = [
        ("自动路由", "不用记数据表和入口，直接用业务语言提问"),
        ("可追溯", "答案保留稳定 ID、来源、时间与证据片段"),
        ("可继续追问", "上下文携带实体，支持从教授跳到论文、专利或企业"),
    ]
    for idx, (title, note) in enumerate(bullets):
        x = 0.88 + idx * 2.34
        add_shape(slide, MSO_SHAPE.OVAL, x, 5.22, 0.13, 0.13, fill=TEAL, line=TEAL)
        add_text(slide, title, x + 0.22, 5.13, 1.62, 0.2, size=9.5, color=SURFACE, bold=True)
        add_text(slide, note, x + 0.22, 5.4, 1.96, 0.39, size=7.7, color=RGBColor(208, 222, 228))

    add_text(slide, "真实系统运行界面", 8.45, 1.82, 3.85, 0.28, size=12.2, color=NAVY_DARK, bold=True)
    add_text(slide, "同一套数据底座，同时服务运营与检索。", 8.45, 2.12, 3.75, 0.24, size=9.6, color=MUTED)
    screenshot_card(slide, home, 8.45, 2.53, 4.28, 2.04, "运营总览 · 数据质量动态", accent=AMBER)
    screenshot_card(slide, professors, 8.45, 4.78, 4.28, 2.04, "教授数据资产 · 条件筛选", accent=TEAL)

    add_text(slide, "深圳科创数据平台", 0.6, 7.14, 3.0, 0.18, size=8, color=MUTED, bold=True)
    add_text(slide, "把分散的科创数据，变成可检索、可解释、可继续追问的答案。", 5.2, 7.1, 7.5, 0.22, size=8.6, color=TEAL, bold=True, align=PP_ALIGN.RIGHT)

    # Slide 2: engineering architecture and operational proof.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CANVAS
    page_header(
        slide,
        "Engineering Foundation",
        "四域数据底座 + Canonical-v2 服务层，支撑持续演进",
        "物理上各域独立建设，逻辑上统一身份、证据、关系与服务契约，工程能力可以持续复用。",
        2,
    )

    add_text(slide, "分层架构", 0.6, 1.82, 2.0, 0.25, size=12.2, color=NAVY_DARK, bold=True)
    add_text(slide, "每一层有清晰边界，替换模型或检索组件不会影响上层业务入口。", 0.6, 2.12, 7.8, 0.24, size=9.6, color=MUTED)

    layer_band(
        slide,
        2.52,
        1.02,
        "交互与运营层",
        TEAL,
        [("用户对话", "自然语言入口"), ("运营总览", "数据质量动态"), ("导入任务", "采集 / 发布"), ("质量 / 审核", "问题闭环")],
    )
    layer_band(
        slide,
        3.84,
        1.28,
        "Canonical-v2 服务层",
        NAVY,
        [("统一身份", "稳定 ID"), ("领域投影", "四域契约"), ("关系投影", "跨域关联"), ("证据落地", "source + time"), ("流式会话", "多轮上下文")],
    )
    layer_band(
        slide,
        5.38,
        1.05,
        "数据与检索层",
        AMBER,
        [("PostgreSQL", "事实与关系"), ("Milvus", "向量召回"), ("Web Search", "实时补充")],
    )
    add_arrow(slide, 4.58, 3.58, 4.58, 3.81, color=TEAL, width=1.4)
    add_arrow(slide, 4.58, 5.13, 4.58, 5.35, color=AMBER, width=1.4)

    rounded(slide, 8.86, 1.82, 3.88, 4.61, fill=NAVY_DARK, line=NAVY_DARK)
    add_text(slide, "工程能力证据", 9.12, 2.08, 2.5, 0.26, size=13, color=SURFACE, bold=True)
    add_text(slide, "不是一条 demo 链路，而是一套可运营的服务底座。", 9.12, 2.4, 3.12, 0.42, size=8.8, color=RGBColor(190, 211, 219))
    evidence_item(slide, 1, "统一合同", "四域对外字段一致，底层 schema 可独立演进。", 2.95, color=TEAL)
    evidence_item(slide, 2, "可追溯", "稳定 ID + evidence，回答可以回到来源。", 3.65, color=TEAL)
    evidence_item(slide, 3, "可迭代", "质量门与 pipeline issues，支持持续修正。", 4.35, color=AMBER)
    evidence_item(slide, 4, "可关联", "关系投影支撑教授、企业、论文、专利跨域跳转。", 5.05, color=TEAL)
    evidence_item(slide, 5, "可验证", "回放 / replay 验证服务行为，降低上线风险。", 5.75, color=AMBER)

    add_text(slide, "运行证据", 0.6, 6.68, 1.4, 0.2, size=8.8, color=MUTED, bold=True)
    slide.shapes.add_picture(str(home), inch(1.82), inch(6.58), width=inch(1.95), height=inch(0.75))
    add_text(slide, "可运营", 3.88, 6.72, 0.85, 0.18, size=8.5, color=AMBER, bold=True)
    slide.shapes.add_picture(str(professors), inch(5.0), inch(6.58), width=inch(1.95), height=inch(0.75))
    add_text(slide, "可追溯", 7.06, 6.72, 0.85, 0.18, size=8.5, color=TEAL, bold=True)
    add_text(slide, "从数据采集、质量控制到服务发布，形成可回看、可修正、可复用的工程闭环。", 8.95, 6.72, 3.75, 0.3, size=8.5, color=NAVY_DARK, bold=True, align=PP_ALIGN.RIGHT)

    add_text(slide, "国际先进技术应用推进中心（深圳） · 客户介绍稿", 0.6, 7.14, 4.8, 0.18, size=8, color=MUTED)
    add_text(slide, "Architecture that stays inspectable.", 9.35, 7.14, 3.4, 0.18, size=8, color=TEAL, bold=True, align=PP_ALIGN.RIGHT, font=FONT_LATIN)

    prs.save(OUT_FILE)
    print(f"wrote {OUT_FILE}")


if __name__ == "__main__":
    build_deck()
