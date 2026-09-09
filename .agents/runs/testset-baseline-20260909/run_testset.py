#!/usr/bin/env python3
"""Run the workbook test set (docs/测试集答案.xlsx) against a live /api/chat/stream.

Three-layer judgment harness (harden-serving-test-harness, GAP-10):
per turn — entity coverage / stance consistency / completeness >= ratio,
plus a provenance layer (local citations, patent ids, web-citation hygiene).
A turn passes only when every applicable layer passes. Keyword hits alone
no longer pass a turn.

Usage:
  python run_testset.py --base-url http://127.0.0.1:18188 --out results.json
  python run_testset.py --offline results-ds-flash-systemd.json   # re-score archived run
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
import urllib.request
from http.cookiejar import CookieJar
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from anchors import ANCHORS, CITATION_FORBIDDEN_PATTERNS  # noqa: E402

REPO = Path("/home/longxiang/MiroThinker")
XLSX = REPO / "docs/测试集答案.xlsx"
PATENT_RX = r"CN\d{9,}[A-Z]?"
_SENTENCE_SPLIT = re.compile(r"[。!?\n]")


def load_sessions() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    sessions: list[dict] = []
    cur: dict | None = None
    for row in ws.iter_rows(min_row=1, values_only=True):
        q, a, k = (list(row) + [None, None, None])[:3]
        q = (str(q) if q else "").strip()
        a = (str(a) if a else "").strip()
        k = (str(k) if k else "").strip()
        if not q:
            continue
        m = re.fullmatch(r"问题(\d+)", q)
        if m:
            cur = {"group": int(m.group(1)), "turns": []}
            sessions.append(cur)
            continue
        if cur is None:
            continue
        cur["turns"].append({"query": q, "expected": a, "key_point": k})
    return sessions


def post_turn(base_url: str, jar: CookieJar, query: str, timeout: int = 180) -> dict:
    body = json.dumps({"query": query}).encode("utf-8")
    req = urllib.request.Request(
        f"{base_url}/api/chat/stream",
        data=body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    started = time.time()
    raw = b""
    try:
        with opener.open(req, timeout=timeout) as resp:
            for chunk in resp:
                raw += chunk
    except Exception as exc:  # noqa: BLE001 - harness records, never raises
        return {"error": f"{type(exc).__name__}: {exc}", "elapsed": time.time() - started,
                "events": {}, "answer": {}, "raw": ""}
    text = raw.decode("utf-8", errors="replace")
    events: dict[str, int] = {}
    answer: dict = {}
    current = None
    for line in text.splitlines():
        if line.startswith("event: "):
            current = line[7:].strip()
            events[current] = events.get(current, 0) + 1
        elif line.startswith("data: {") and current == "answer":
            try:
                answer = json.loads(line[6:])
            except json.JSONDecodeError:
                pass
    return {"error": None, "elapsed": round(time.time() - started, 1),
            "events": events, "answer": answer, "raw_len": len(text)}


def _hit(term: str, low: str) -> bool:
    return term.casefold() in low


def _kp_hit(kp, low: str) -> bool:
    if isinstance(kp, (list, tuple)):
        return any(_hit(str(a), low) for a in kp)
    return _hit(str(kp), low)


def score_turn(text: str, anchor: dict, *, citations_local: int | None = None,
               citation_texts: list[str] | None = None) -> dict:
    """Three-layer + provenance scoring. Layers without anchors are skipped."""
    low = text.casefold()
    layers: dict[str, dict] = {}

    # Layer 1: entity coverage
    fails: list[str] = []
    if anchor.get("entities") or anchor.get("entities_any") or anchor.get("entities_pool") or anchor.get("forbidden"):
        for term in anchor.get("entities", ()):
            if not _hit(term, low):
                fails.append(f"missing:{term}")
        any_terms = anchor.get("entities_any")
        if any_terms and not any(_hit(t, low) for t in any_terms):
            fails.append(f"missing_any:{'/'.join(any_terms)}")
        pool = anchor.get("entities_pool")
        if pool:
            hits = sum(1 for t in pool if _hit(t, low))
            need = int(anchor.get("entities_min", 1))
            if hits < need:
                fails.append(f"pool:{hits}<{need}")
        for term in anchor.get("forbidden", ()):
            if _hit(term, low):
                fails.append(f"forbidden:{term}")
        layers["entity"] = {"pass": not fails, "fails": fails}

    # Layer 2: stance consistency (sentence-scoped regexes + fact anchors)
    fails = []
    stance_pats = anchor.get("stance_forbid") or ()
    req_if = anchor.get("require_if")
    if stance_pats or req_if:
        sentences = [s for s in _SENTENCE_SPLIT.split(text) if s.strip()]
        for pat in stance_pats:
            for sent in sentences:
                if re.search(pat, sent):
                    fails.append(f"stance:{pat[:40]} @ {sent.strip()[:40]}")
                    break
        if req_if and _hit(req_if["trigger"], low):
            if not any(_hit(t, low) for t in req_if["required_any"]):
                fails.append(f"fact:{req_if['trigger']} without {'/'.join(req_if['required_any'])}")
        layers["stance"] = {"pass": not fails, "fails": fails}

    # Layer 3: completeness (key-point coverage ratio)
    kps = anchor.get("key_points") or ()
    if kps:
        hits = sum(1 for kp in kps if _kp_hit(kp, low))
        ratio = float(anchor.get("key_points_ratio", 0.8))
        ok = hits / len(kps) >= ratio
        layers["completeness"] = {
            "pass": ok, "hits": hits, "total": len(kps), "ratio": ratio,
            "fails": [] if ok else [f"coverage:{hits}/{len(kps)}<{ratio}"],
        }

    # Layer 4: provenance (local citations / patent ids / web hygiene)
    fails = []
    prov_applicable = False
    if anchor.get("local_citations_min") and citations_local is not None:
        prov_applicable = True
        need = int(anchor["local_citations_min"])
        if citations_local < need:
            fails.append(f"local_citations:{citations_local}<{need}")
    if anchor.get("patent_ids_min"):
        prov_applicable = True
        hits = len(re.findall(PATENT_RX, text))
        if hits < int(anchor["patent_ids_min"]):
            fails.append(f"patent_ids:{hits}<{anchor['patent_ids_min']}")
    if citation_texts:
        prov_applicable = True
        for pat in CITATION_FORBIDDEN_PATTERNS:
            for ct in citation_texts:
                if re.search(pat, ct, re.IGNORECASE):
                    fails.append(f"web_boilerplate:{pat} @ {ct[:40]}")
                    break
    if prov_applicable:
        layers["provenance"] = {"pass": not fails, "fails": fails}

    overall = all(l["pass"] for l in layers.values()) if layers else False
    return {"pass": overall, "layers": layers, "note": anchor.get("note", ""),
            "human_verified": bool(anchor.get("human_verified"))}


def _anchor_for(group: int, turn_idx: int) -> dict:
    return ANCHORS.get(f"g{group}-t{turn_idx}", {"note": "(no anchor)"})


def _citation_texts(citations: list) -> list[str]:
    out = []
    for c in citations:
        if not isinstance(c, dict):
            continue
        parts = [str(c.get(k) or "") for k in ("title", "snippet", "locator", "url", "source_locator")]
        out.append(" | ".join(p for p in parts if p))
    return out


def run_live(args) -> list[dict]:
    sessions = load_sessions()
    only = {int(x) for x in args.only.split(",") if x.strip()}
    results = []
    for sess in sessions:
        g = sess["group"]
        if only and g not in only:
            continue
        jar = CookieJar()
        turn_results = []
        for idx, turn in enumerate(sess["turns"], 1):
            res = post_turn(args.base_url, jar, turn["query"], args.timeout)
            answer = res.get("answer") or {}
            text = answer.get("answer_text") or ""
            cits = answer.get("citations") or []
            local = [c for c in cits if str(c.get("type")) != "web"]
            web = [c for c in cits if str(c.get("type")) == "web"]
            anchor = _anchor_for(g, idx)
            scored = score_turn(text, anchor, citations_local=len(local),
                                citation_texts=_citation_texts(web)) if text else {
                "pass": False, "layers": {}, "note": anchor.get("note", ""),
                "fails": ["empty_answer"]}
            turn_results.append({
                "turn": idx, "query": turn["query"], "key_point": turn["key_point"],
                "error": res.get("error"), "elapsed": res.get("elapsed"),
                "query_type": answer.get("query_type"), "answer_style": answer.get("answer_style"),
                "answer_text": text, "citations_total": len(cits),
                "citations_local": len(local), "citations_web": len(web),
                "events": res.get("events"), "scored": scored,
            })
            _print_turn(g, idx, res, answer, len(cits), len(local), len(web), turn, scored)
        results.append({"group": g, "turns": turn_results,
                        "passed": sum(1 for t in turn_results if t["scored"]["pass"]),
                        "total": len(turn_results)})
    return results


def run_offline(path: Path) -> list[dict]:
    archived = json.loads(path.read_text(encoding="utf-8"))
    results = []
    for grp in archived:
        g = grp["group"]
        turn_results = []
        for t in grp["turns"]:
            idx = t["turn"]
            text = t.get("answer_text") or ""
            anchor = _anchor_for(g, idx)
            scored = score_turn(text, anchor,
                                citations_local=t.get("citations_local")) if text else {
                "pass": False, "layers": {}, "note": anchor.get("note", ""),
                "fails": ["empty_answer"]}
            nt = dict(t)
            nt["scored"] = scored
            turn_results.append(nt)
            _print_turn(g, idx, {"elapsed": t.get("elapsed"), "error": t.get("error")},
                        {"query_type": t.get("query_type")},
                        t.get("citations_total", 0), t.get("citations_local", 0),
                        t.get("citations_web", 0), t, scored)
        results.append({"group": g, "turns": turn_results,
                        "passed": sum(1 for t in turn_results if t["scored"]["pass"]),
                        "total": len(turn_results)})
    return results


def _print_turn(g, idx, res, answer, cit_total, cit_local, cit_web, turn, scored) -> None:
    status = "PASS" if scored["pass"] else "FAIL"
    print(f"[g{g}-t{idx}] {status} {res.get('elapsed')}s {answer.get('query_type')} "
          f"cit={cit_total}(L{cit_local}/W{cit_web}) :: {turn['query'][:36]}", flush=True)
    for lname, layer in (scored.get("layers") or {}).items():
        if not layer["pass"]:
            print(f"        {lname}: {layer.get('fails')}", flush=True)
    if scored.get("fails"):
        print(f"        fails: {scored['fails']}", flush=True)


def print_summary(results: list[dict]) -> None:
    tp = sum(r["passed"] for r in results)
    tt = sum(r["total"] for r in results)
    print(f"\n==== TOTAL {tp}/{tt} turns pass (three-layer) ====")
    layer_names = ("entity", "stance", "completeness", "provenance")
    for lname in layer_names:
        app = pas = 0
        for r in results:
            for t in r["turns"]:
                layer = (t["scored"].get("layers") or {}).get(lname)
                if layer is not None:
                    app += 1
                    pas += 1 if layer["pass"] else 0
        if app:
            print(f"  layer {lname}: {pas}/{app}")
    for r in results:
        print(f"  g{r['group']}: {r['passed']}/{r['total']}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base-url", default="http://127.0.0.1:18188")
    ap.add_argument("--out", default=str(Path(__file__).parent / "results.json"))
    ap.add_argument("--only", default="", help="comma-separated group numbers")
    ap.add_argument("--timeout", type=int, default=180)
    ap.add_argument("--offline", default="", help="re-score an archived results JSON (no HTTP)")
    args = ap.parse_args()

    if args.offline:
        results = run_offline(Path(args.offline))
    else:
        results = run_live(args)
        out = Path(args.out)
        out.write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"results -> {out}")
    print_summary(results)
    return 0


if __name__ == "__main__":
    sys.exit(main())
